"""Read-only check of the delivered XLSX caches, source rows, and native chart."""
import collections
import hashlib
import json
import math
import pathlib
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

PROJECT=pathlib.Path(__file__).resolve().parents[1]
OUT=PROJECT/'outputs/01a06f83-a122-76c3-b29c-df4417405f3b'
FILE=OUT/'港股通469股_20策略_十年回测与评分_2026-09-04.xlsx'
data=json.loads((PROJECT/'data/connect-10y-2026-09-06/workbook-data.json').read_text(encoding='utf-8'))
wb=openpyxl.load_workbook(FILE,read_only=True,data_only=True)
assert wb.sheetnames==['总览','十年策略评分','扩展样本评分','全量回测','收益矩阵','股票与数据','方法与参数']

def check(actual, expected, label):
    if expected is None:
        assert actual in (None,''),(label,actual,expected)
    elif isinstance(expected,(int,float)):
        assert isinstance(actual,(int,float)) and math.isclose(actual,expected,rel_tol=1e-9,abs_tol=1e-9),(label,actual,expected)
    else:
        assert actual==expected,(label,actual,expected)

raw=list(wb['全量回测'].iter_rows(min_row=7,max_row=9386,max_col=29,values_only=True))
assert len(raw)==9380
for i,(a,e) in enumerate(zip(raw,data['rows'])):
    for j,key in [(0,'code'),(1,'name'),(2,'strategy_id'),(6,'years'),(7,'bars'),(8,'full10'),(9,'extended'),(10,'status'),
                  (11,'initial'),(12,'final'),(13,'total_return'),(14,'annual_return'),(15,'drawdown'),(16,'sharpe'),
                  (17,'trades'),(18,'win_rate'),(19,'profit_factor'),(20,'excess_annual'),(22,'norm_return'),
                  (23,'norm_drawdown'),(24,'norm_sharpe'),(25,'score')]:
        check(a[j],e[key],f'row {i+7} {key}')
assert len({r[0] for r in raw})==469
assert collections.Counter(r[0] for r in raw).most_common(1)[0][1]==20
assert sum(r[10]=='数据异常' for r in raw)==100

for cohort,sheet in [('full10','十年策略评分'),('extended','扩展样本评分')]:
    rows=list(wb[sheet].iter_rows(min_row=7,max_row=26,min_col=3,max_col=21,values_only=True))
    for actual,expected in zip(rows,data['rankings'][cohort]):
        fields=['rank','id','name','score','valid','samples','coverage','norm_return','norm_drawdown','norm_sharpe',
                'mean_annual','median_annual','mean_total','median_total','mean_drawdown','mean_sharpe','profit_rate','outperform_rate','few_trades']
        for j,key in enumerate(fields):
            e=None if expected['id']=='S01' and key=='outperform_rate' else expected[key]
            check(actual[j],e,f'{sheet} {expected["id"]} {key}')

matrix=list(wb['收益矩阵'].iter_rows(min_row=7,max_row=475,min_col=8,max_col=27,values_only=True))
for i,row in enumerate(matrix):
    for j,value in enumerate(row):check(value,data['rows'][i*20+j]['total_return'],f'matrix {i} {j}')
for sheet in wb:
    for row in sheet.iter_rows():
        for cell in row:
            assert cell.data_type!='e',(sheet.title,cell.coordinate,cell.value)
wb.close()
with zipfile.ZipFile(FILE) as z:
    chart_files=[n for n in z.namelist() if n.startswith('xl/') and '/charts/chart' in n and n.endswith('.xml')]
    assert len(chart_files)==1,chart_files
    assert b'barChart' in z.read(chart_files[0])
    all_formulas=sum(sum(1 for _ in ET.fromstring(z.read(n)).iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f'))
                     for n in z.namelist() if n.startswith('xl/worksheets/sheet') and n.endswith('.xml'))
    assert all_formulas>80000,all_formulas
result=dict(status='passed',stocks=469,strategies=20,records=9380,successful_stocks=464,failed_stocks=5,
            full10=275,extended=368,formula_cells=all_formulas,native_charts=len(chart_files),
            sha256=hashlib.sha256(FILE.read_bytes()).hexdigest(),bytes=FILE.stat().st_size)
(OUT/'xlsx-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
